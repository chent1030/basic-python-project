# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = '/Users/csai/project/basic-project/CPS智能巡检系统PRD.docx'

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)

def cell_text(cell, text, bold=False, color=None, size=8.5):
    cell.text = ''
    p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
    r = p.add_run(str(text)); r.bold = bold; r.font.size = Pt(size)
    if color: r.font.color.rgb = RGBColor(*color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def add_table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers)); t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell_text(t.rows[0].cells[i], h, True, (255,255,255), 9); shade(t.rows[0].cells[i], '1F6FEB')
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row): cell_text(cells[i], value)
    doc.add_paragraph()

def bullets(doc, items):
    for item in items: doc.add_paragraph(item, style='List Bullet')

doc = Document(); sec = doc.sections[0]
sec.top_margin = Inches(.65); sec.bottom_margin = Inches(.65); sec.left_margin = Inches(.75); sec.right_margin = Inches(.75)
for name, size, color in [('Normal',10,(21,32,43)), ('Title',25,(21,32,43)), ('Heading 1',17,(31,111,235)), ('Heading 2',13,(15,118,110))]:
    s = doc.styles[name]; s.font.name = 'Microsoft YaHei'; s._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei'); s.font.size = Pt(size); s.font.color.rgb = RGBColor(*color)

p = doc.add_paragraph(style='Title'); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run('CPS 智能巡检系统产品需求文档')
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; r = p.add_run('版本 1.0  |  2026 年 9 月 9 日  |  产品设计稿'); r.font.size = Pt(10); r.font.color.rgb = RGBColor(102,114,127)
doc.add_paragraph('本 PRD 定义 CPS 巡检场景下的业务闭环、主 Agent 动态调度、人工确认机制、旁路观察与长期记忆能力。核心原则是：AI 提出建议，人工确认或修正下一步动作，系统记录决策并持续优化。')

doc.add_heading('1 文档概述', 1)
add_table(doc, ['项目','内容'], [['产品名称','CPS 智能巡检系统'],['目标用户','PDA 巡检员工、拉线主管、整改负责人、巡检管理员'],['核心场景','产线问题发现、整改跟踪、整改复查、报告汇总与持续优化'],['技术方向','LangChain + LangGraph Runtime + DeepAgents + 多模态模型 + 向量库'],['本版本目标','打通一次巡检从采集到确认、入库和任务生成的可审计闭环']])
doc.add_heading('2 背景与目标', 1); doc.add_heading('2.1 业务背景', 2)
doc.add_paragraph('员工通过 PDA 检查产线问题，现场上传图片并跟踪整改。系统需要统一现场信息、历史巡检记录和整改证据，减少人工汇总成本，提升问题识别、整改判断和后续工作安排的质量。')
doc.add_heading('2.2 产品目标', 2)
bullets(doc, ['将现场图片、拉线改造信息和历史巡检记录沉淀为统一巡检任务。','通过多 Agent 协作完成问题识别、整改判断、历史分析、报告生成和覆盖分析。','每次调用下一步 Agent 前引入人工确认，允许批准、修改、跳过或退回。','将人工调度选择、修改原因和后续结果沉淀为长期记忆，经审核后优化未来流程与决策。','所有最终业务事实、报告入库和任务发布均可追溯、可审计。'])
doc.add_heading('2.3 非目标', 2)
bullets(doc, ['本期不自动替代巡检人员对业务结论的最终确认。','本期不允许观察 Agent 直接修改主 Agent 的 Prompt、规则或线上策略。','统计指标由结构化查询和代码计算，模型负责解释和提出建议。'])

doc.add_heading('3 用户与权限', 1)
add_table(doc, ['角色','主要职责','关键权限'], [['巡检员工','创建任务、采集图片、补充现场信息','提交任务、修改本人未确认内容'],['拉线主管','确认问题与整改状态、补充整改要求','确认、编辑、退回问题和整改结论'],['巡检管理员','确认报告、入库、审核记忆和优化建议','发布报告、审核长期记忆、管理规则'],['系统管理员','配置模型、Agent、工具和权限','配置与审计，不参与业务结论确认']])

doc.add_heading('4 总体业务流程', 1)
doc.add_paragraph('系统以 InspectionCase（巡检任务）为核心对象。主 Agent 不使用预先固定的业务节点顺序，而是根据当前状态和人工反馈持续提出下一步动作。')
add_table(doc, ['阶段','流程','人工参与点','输出'], [['现场采集','创建任务 → PDA 上传问题图片 → 补充位置和备注','无或补充信息','现场问题证据'],['问题分析','主 Agent 调用问题识别 Agent → 生成候选结论','确认、修改、补充或驳回问题','已确认问题'],['整改复查','提交整改后图片 → 主 Agent 调用整改判断 Agent','确认是否执行判断、确认整改结论','最终整改状态'],['报告闭环','调用历史分析、覆盖分析、报告 Agent','确认报告、确认入库','HTML 报告与结构化数据'],['持续优化','观察 Agent 监听链路 → 筛选记忆 → 审核后入库','审核记忆与优化建议','长期记忆与优化策略']])

doc.add_heading('5 主 Agent 动态调度', 1); doc.add_heading('5.1 主 Agent 职责', 2)
bullets(doc, ['理解用户目标、任务类型和当前巡检状态。','读取短期工作记忆、历史案例、规则和人工反馈。','判断当前是否具备执行条件，并提出下一步 Agent 或工具。','在执行下一步前创建人工调度确认请求。','根据人工的批准、改派、跳过、重试或补充指令继续执行。','接收结果，更新任务状态，再次决定下一步，直到完成或转人工。'])
doc.add_heading('5.2 人工调度确认卡', 2)
add_table(doc, ['字段','说明'], [['当前状态','已完成的动作和证据完整度'],['主 Agent 建议','建议调用的 Agent、调用理由、预期输出'],['候选动作','批准、修改 Agent、跳过、重试、退回采集、人工处理'],['人工修改','最终选择的 Agent、补充指令、修改原因'],['确认人','用户、角色、时间'],['后续结果','Agent 输出、是否成功、整改效果、是否产生新记忆']])
doc.add_heading('5.3 调度规则', 2)
bullets(doc, ['低置信度、证据不足、规则冲突或高风险问题必须请求人工确认。','人工选择优先于主 Agent 建议，但不得绕过权限和数据校验。','人工可以选择候选 Agent 之外的允许 Agent，但不能调用无权限工具。','人工选择和修改必须原样记录，不能只保存最终结果。','主 Agent 因结果不完整再次提出下一步时，必须重新进入确认。'])

doc.add_heading('6 Agent 能力定义', 1)
add_table(doc, ['Agent','输入','输出'], [['问题识别 Agent','现场图片、产线/区域信息、相似案例','问题分类、描述、位置、严重程度、整改建议、证据和置信度'],['整改判断 Agent','整改前后图片、原问题、整改要求、规范案例','到位/部分到位/不到位/无法判断、证据、剩余风险和下一步建议'],['历史分析 Agent','历史巡检记录、时间范围、筛选条件','问题频率、区域排行、主管排行、复发率、措施有效性'],['报告 Agent','已确认问题、整改结论、历史分析、覆盖分析','HTML 报告草稿、摘要、重点问题、风险和行动建议'],['覆盖分析 Agent','本次覆盖范围、历史高风险范围、采集证据','未覆盖场景、数据缺口、风险提示和补充巡检建议'],['观察 Agent','全链路事件、人工调度、结果、复发和整改效果','候选长期记忆、失败模式、流程优化建议和策略评估']])

doc.add_heading('7 长期记忆与流程优化', 1); doc.add_heading('7.1 监听范围', 2)
bullets(doc, ['主 Agent 的任务理解、建议动作和决策理由。','人工确认、改派 Agent、跳过、重试、退回及人工补充说明。','专用 Agent 的输入、输出、置信度、失败原因和工具调用。','报告最终内容与人工修改内容。','整改复查结果、后续复发情况和工作安排执行效果。'])
doc.add_heading('7.2 可沉淀的记忆类型', 2)
add_table(doc, ['类型','示例','用途'], [['业务案例记忆','某区域某类问题最终确认与有效整改方式','辅助问题识别和整改建议'],['调度偏好记忆','某类证据不足时优先调用人工复核','优化下一步 Agent 推荐'],['人工修正记忆','模型常将某现象分类为 A，主管修正为 B','降低重复误判'],['失败模式记忆','某类整改图片角度不足导致无法判断','生成采集提示和补拍要求'],['流程优化记忆','报告生成前必须补充某区域数据','完善流程检查和覆盖分析'],['策略评估记忆','某条调用路径后续复发率更低','评估是否推荐该路径']])
doc.add_heading('7.3 记忆写入原则', 2)
bullets(doc, ['观察 Agent 只生成候选记忆，不直接写入生产长期记忆。','候选记忆经过管理员审核后才生效。','记忆必须带场景、证据、来源、置信度、创建时间和适用范围。','重复或相互冲突的记忆进入合并或冲突队列。','长期记忆可以被停用、回滚和追溯。'])

doc.add_heading('8 人工确认与状态管理', 1)
add_table(doc, ['状态','含义'], [['pending_dispatch_confirmation','主 Agent 已提出下一步，等待人工确认'],['approved','人工同意主 Agent 建议'],['modified','人工改派 Agent 或补充调用指令'],['skipped','人工跳过当前建议'],['returned','人工退回重新采集或整改'],['completed','Agent 执行成功并返回结果'],['needs_human','Agent 无法完成或结果冲突，转人工处理']])
doc.add_paragraph('任何写入历史巡检库、发布报告、创建整改任务或发送通知的动作，都必须支持幂等，并保留操作人和事件记录。')

doc.add_heading('9 数据对象', 1)
add_table(doc, ['对象','关键字段'], [['InspectionCase','inspection_id、line_info、status、created_by、thread_id'],['AgentDispatch','dispatch_id、recommended_agent、selected_agent、decision、reason、operator_id'],['AgentRun','run_id、agent_name、input_ref、output、confidence、error'],['HumanConfirmation','confirmation_id、stage、original_payload、edited_payload、comment'],['MemoryCandidate','memory_id、type、content、evidence、source_events、review_status'],['LongTermMemory','scope、condition、knowledge、effective、version、reviewer'],['Report','report_id、draft_data、confirmed_data、html_ref、archive_status']])

doc.add_heading('10 产品功能需求', 1)
add_table(doc, ['编号','需求','优先级','验收标准'], [['FR-01','创建巡检任务并维护拉线改造信息','P0','可保存产线、区域、主管和任务状态'],['FR-02','上传并管理现场图片及整改前后证据','P0','图片可关联任务、问题和整改记录'],['FR-03','主 Agent 动态推荐下一步 Agent','P0','推荐包含 Agent、理由、前置条件和预期输出'],['FR-04','人工确认下一步 Agent','P0','可批准、改派、跳过、重试、退回或补充指令'],['FR-05','Agent 结构化输出和结果回传','P0','结果可被主 Agent 读取并用于下一次决策'],['FR-06','生成并编辑巡检报告','P0','报告可预览、修改、确认并导出 HTML'],['FR-07','报告转结构化历史数据','P0','确认后才入库，支持幂等和审计'],['FR-08','观察 Agent 监听全链路','P1','可查看事件链和候选优化建议'],['FR-09','长期记忆审核与版本管理','P1','可审核、启用、停用、回滚记忆'],['FR-10','历史分析与覆盖分析','P1','可输出频率、区域、主管、复发和覆盖缺口']])

doc.add_heading('11 非功能需求', 1)
bullets(doc, ['可追溯：所有 Agent 调用、人工决策、数据写入和记忆变更均有事件日志。','可靠性：主 Agent 中断后可从最近一次确认点恢复；重复恢复不重复写库。','安全性：Agent 和工具按角色授权，图片、历史记录和长期记忆隔离访问。','可解释性：每个建议包含证据、理由、置信度和数据来源。','可运营性：支持模型、Agent、工具、记忆和策略的版本化管理。','可控性：线上长期记忆和策略变更必须经过审核，并支持回滚。'])
doc.add_heading('12 MVP 范围与迭代计划', 1)
add_table(doc, ['阶段','范围'], [['MVP','任务创建、图片采集、问题识别、人工确认调度、整改判断、人工确认、报告确认、结构化入库'],['第二阶段','历史分析、覆盖分析、工作安排、观察 Agent、候选记忆审核'],['第三阶段','记忆效果评估、策略版本管理、自动化复查任务、调度路径优化、质量看板']])
doc.add_heading('13 关键指标', 1)
add_table(doc, ['指标','定义'], [['问题识别人工修改率','人工修改问题识别结果的任务数 / 触发问题识别的任务数'],['Agent 改派率','人工修改主 Agent 推荐 Agent 的调度次数 / 总调度次数'],['人工确认通过率','直接批准主 Agent 建议的次数 / 总确认次数'],['整改判断准确率','人工最终结论与 Agent 结论一致的比例'],['记忆采纳率','审核后生效的候选记忆数 / 候选记忆总数'],['流程优化收益','策略生效后任务时长、复发率或人工操作次数的变化']])
doc.add_heading('14 风险与控制措施', 1)
add_table(doc, ['风险','控制措施'], [['主 Agent 错误选择 Agent','强制人工调度确认、权限校验和候选 Agent 白名单'],['错误记忆污染后续决策','候选记忆审核、来源证据、版本控制和回滚'],['人工确认过多影响效率','按风险和置信度分级，高风险强确认，低风险可配置批量确认'],['统计被模型误解','统计由 SQL/代码计算，Agent 只解释并引用指标'],['重复写库或重复发通知','幂等键、事务状态和事件去重']])
doc.add_heading('15 待确认事项', 1)
bullets(doc, ['人工确认是每次调用都确认，还是按风险等级支持可配置豁免。','长期记忆的审核角色、审核时限和失效周期。','可被主 Agent 调用的 Agent 与工具白名单。','历史记录的观察窗口、复发定义和整改有效性计算口径。','HTML 报告模板、品牌规范和导出方式。'])

footer = sec.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER; footer.add_run('CPS 智能巡检系统 PRD | 版本 1.0').font.size = Pt(8)
doc.save(OUT); print(OUT)
